from io import BytesIO

import pandas as pd

from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)

from openpyxl.utils import get_column_letter


# ============================================================
# OUTILS DE MISE EN FORME
# ============================================================

def mettre_en_forme_feuille(ws):
    """
    Applique une mise en forme simple et professionnelle
    à une feuille Excel.
    """

    # Couleur sobre pour les en-têtes
    remplissage_entete = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    police_entete = Font(
        bold=True,
        color="1F1F1F"
    )

    bordure_fine = Border(
        bottom=Side(
            style="thin",
            color="B7B7B7"
        )
    )

    # --------------------------------------------------------
    # En-têtes
    # --------------------------------------------------------

    for cellule in ws[1]:
        cellule.fill = remplissage_entete
        cellule.font = police_entete
        cellule.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
        cellule.border = bordure_fine

    # --------------------------------------------------------
    # Figer la première ligne
    # --------------------------------------------------------

    ws.freeze_panes = "A2"

    # --------------------------------------------------------
    # Activer les filtres
    # --------------------------------------------------------

    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions

    # --------------------------------------------------------
    # Hauteur de l'en-tête
    # --------------------------------------------------------

    ws.row_dimensions[1].height = 35

    # --------------------------------------------------------
    # Largeur automatique des colonnes
    # --------------------------------------------------------

    for numero_colonne, colonne in enumerate(
        ws.iter_cols(),
        start=1
    ):
        longueur_max = 0

        for cellule in colonne:
            if cellule.value is not None:
                longueur = len(str(cellule.value))
                longueur_max = max(
                    longueur_max,
                    longueur
                )

        # On limite volontairement les colonnes très longues
        largeur = min(
            max(longueur_max + 2, 12),
            45
        )

        ws.column_dimensions[
            get_column_letter(numero_colonne)
        ].width = largeur

    # --------------------------------------------------------
    # Alignement du contenu
    # --------------------------------------------------------

    for ligne in ws.iter_rows(
        min_row=2
    ):
        for cellule in ligne:
            cellule.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


# ============================================================
# CONSTRUCTION DE LA SYNTHÈSE
# ============================================================

def creer_synthese(df):
    """
    Construit une synthèse des rejets et des décisions
    prises par le prototype.
    """

    if "Résultat" not in df.columns:
        return pd.DataFrame()

    lignes_synthese = []

    for resultat, groupe in df.groupby(
        "Résultat",
        dropna=False
    ):

        nombre_total = len(groupe)

        nombre_regle = int(
            (
                groupe["Mode de décision"]
                == "Règle métier"
            ).sum()
        )

        nombre_humain = int(
            (
                groupe["Mode de décision"]
                == "Traitement humain"
            ).sum()
        )

        traitements = (
            groupe["Traitement proposé"]
            .dropna()
            .astype(str)
            .unique()
            .tolist()
        )

        if traitements:
            traitement_propose = " / ".join(
                traitements
            )
        else:
            traitement_propose = (
                "Aucun traitement automatique"
            )

        taux_couverture = (
            nombre_regle / nombre_total * 100
            if nombre_total > 0
            else 0
        )

        lignes_synthese.append(
            {
                "Résultat SITR": resultat,
                "Nombre de cas": nombre_total,
                "Traitement proposé": traitement_propose,
                "Traités par règle métier": nombre_regle,
                "À analyser humainement": nombre_humain,
                "Taux de couverture (%)": round(
                    taux_couverture,
                    1
                ),
            }
        )

    synthese = pd.DataFrame(
        lignes_synthese
    )

    synthese = synthese.sort_values(
        by="Nombre de cas",
        ascending=False
    )

    return synthese


# ============================================================
# EXPORT EXCEL
# ============================================================

def dataframe_to_excel(df):
    """
    Génère le fichier Excel final du prototype.

    Le classeur contient trois feuilles :

    1. ACT Métier enrichi
       Toutes les données d'origine et les propositions
       du prototype.

    2. Synthèse des rejets
       Vue agrégée des types de rejets et traitements.

    3. Cas à analyser
       Cas ne disposant pas d'une règle suffisamment
       fiable dans le périmètre actuel du PoC.
    """

    output = BytesIO()

    # --------------------------------------------------------
    # Préparation des données
    # --------------------------------------------------------

    donnees_enrichies = df.copy()

    synthese = creer_synthese(
        donnees_enrichies
    )

    cas_a_analyser = (
        donnees_enrichies[
            donnees_enrichies[
                "Mode de décision"
            ] == "Traitement humain"
        ]
        .copy()
    )

    # --------------------------------------------------------
    # Création du classeur
    # --------------------------------------------------------

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        # Feuille 1
        donnees_enrichies.to_excel(
            writer,
            index=False,
            sheet_name="ACT Métier enrichi"
        )

        # Feuille 2
        synthese.to_excel(
            writer,
            index=False,
            sheet_name="Synthèse des rejets"
        )

        # Feuille 3
        cas_a_analyser.to_excel(
            writer,
            index=False,
            sheet_name="Cas à analyser"
        )

        # ----------------------------------------------------
        # Mise en forme
        # ----------------------------------------------------

        for nom_feuille in writer.book.sheetnames:

            feuille = writer.book[
                nom_feuille
            ]

            mettre_en_forme_feuille(
                feuille
            )

    output.seek(0)

    return output